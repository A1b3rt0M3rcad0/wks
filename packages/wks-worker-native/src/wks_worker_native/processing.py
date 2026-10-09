import csv
import io
import json
import subprocess
import zipfile
from importlib.metadata import version
from pathlib import Path

from PIL import Image
from wks_core.domain.models import Block, Error, ExtractedAsset, ExtractedRepresentation
from wks_core.storage.content import validate_zip
from wks_worker.processing import text_blocks


class TextProcessor:
    def extract(self, path, mime):
        try:
            value = path.read_text(encoding="utf-8-sig")
            if mime == "application/json":
                json.loads(value)
            elif mime in {"application/yaml", "text/yaml"}:
                import yaml

                yaml.safe_load(value)
        except (ValueError, UnicodeError, RecursionError) as exc:
            return ExtractedRepresentation(
                availability="metadata_only",
                warnings=[f"invalid_text:{type(exc).__name__}"],
                coverage={"text": {"state": "failed"}},
            )
        if mime == "text/csv":
            rows = list(csv.reader(io.StringIO(value)))
            blocks = []
            # Chunk rows instead of placing arbitrary CSV content in a single response.
            for start in range(0, len(rows), 20):
                chunk = rows[start : start + 20]
                blocks.append(
                    Block(
                        "table",
                        "\n".join(", ".join(r) for r in chunk),
                        {"kind": "row_range", "start": start + 1, "end": start + len(chunk)},
                        data={"rows": chunk},
                    )
                )
        else:
            blocks = text_blocks(value)
        return ExtractedRepresentation(
            blocks=blocks, coverage={"text": {"state": "complete"}}, processor_name="text"
        )


class ImageProcessor:
    def __init__(self, settings):
        self.settings = settings

    def ocr(self, path):
        try:
            result = subprocess.run(
                ["tesseract", str(path), "stdout", "-l", self.settings.ocr_language, "--psm", "3"],
                capture_output=True,
                timeout=self.settings.processing_timeout_seconds,
                check=True,
            )
            return result.stdout.decode().strip(), None
        except (OSError, subprocess.SubprocessError):
            return "", "ocr_provider_unavailable"

    def extract(self, path, mime):
        Image.MAX_IMAGE_PIXELS = 40_000_000
        with Image.open(path) as img:
            img.verify()
        data = path.read_bytes()
        a = ExtractedAsset(data, mime, "image", {"kind": "image", "value": "original"})
        b = Block("picture_ref", locator=a.locator, asset_refs=[f"wks://assets/{a.id}"])
        result = ExtractedRepresentation(
            blocks=[b],
            assets=[a],
            availability="metadata_only",
            processor_name="image",
            processor_version=version("pillow"),
            coverage={"text": {"state": "pending"}, "visual": {"state": "not_interpreted"}},
        )
        if self.settings.ocr_enabled:
            value, warning = self.ocr(path)
            result.blocks += text_blocks(value, a.locator, "ocr")
            for block in result.blocks:
                block.asset_refs = [f"wks://assets/{a.id}"]
            result.coverage["ocr"] = {"state": "failed" if warning else "complete"}
            result.coverage["text"] = {"state": "partial" if warning else "complete"}
            result.availability = "text_partial" if warning else "text_ready"
            if warning:
                result.warnings.append(warning)
        return result


class PDFProcessor:
    """Native text first; preserve rendered pages and embedded image bytes; selective OCR."""

    def __init__(self, settings):
        self.settings = settings

    def extract(self, path, mime):
        import pypdfium2
        from pypdf import PdfReader

        pdf = PdfReader(path)
        if pdf.is_encrypted:
            raise Error("upload.content_invalid", "Encrypted PDF cannot be processed")
        renderer = pypdfium2.PdfDocument(str(path))
        out = ExtractedRepresentation(
            processor_name="native-pdf", processor_version=version("pypdf")
        )
        pending, completed = 0, 0
        try:
            for page_no, page in enumerate(pdf.pages[: self.settings.process_max_pages], 1):
                loc = {"kind": "page", "value": str(page_no)}
                out.blocks.append(Block("heading", f"Page {page_no}", loc, data={"level": 2}))
                try:
                    value = page.extract_text() or ""
                except Exception:
                    value = ""
                    out.warnings.append(f"page_{page_no}_native_extraction_failed")
                import math

                size = renderer[page_no - 1].get_size()
                scale = min(2, math.sqrt(16_000_000 / max(1, size[0] * size[1])))
                page_image = renderer[page_no - 1].render(scale=scale).to_pil()
                bio = io.BytesIO()
                page_image.save(bio, format="PNG")
                page_asset = ExtractedAsset(bio.getvalue(), "image/png", "page", loc)
                out.assets.append(page_asset)
                page_refs = [f"wks://assets/{page_asset.id}"]
                try:
                    for picture in page.images:
                        img = Image.open(io.BytesIO(picture.data))
                        media_type = Image.MIME.get(img.format, "application/octet-stream")
                        a = ExtractedAsset(picture.data, media_type, "figure", loc)
                        out.assets.append(a)
                        page_refs.append(f"wks://assets/{a.id}")
                except Exception:
                    out.warnings.append(f"page_{page_no}_image_preserved_as_page")
                origin = "native_text"
                if not value.strip() and self.settings.ocr_enabled:
                    from tempfile import NamedTemporaryFile

                    with NamedTemporaryFile(suffix=".png") as f:
                        f.write(page_asset.content)
                        f.flush()
                        value, warning = ImageProcessor(self.settings).ocr(Path(f.name))
                    origin = "ocr"
                    if warning:
                        out.warnings.append(f"page_{page_no}_{warning}")
                if value.strip():
                    completed += 1
                    blocks = text_blocks(value, loc, origin)
                    for block in blocks:
                        block.asset_refs = page_refs
                    out.blocks += blocks
                else:
                    pending += 1
                out.blocks.append(Block("picture_ref", locator=loc, asset_refs=page_refs))
            pending += max(0, len(pdf.pages) - self.settings.process_max_pages)
            out.coverage = {
                "text": {
                    "state": "partial" if pending else "complete",
                    "processed_pages": completed,
                    "total_pages": len(pdf.pages),
                    "pending_pages": pending,
                },
                "visual": {"state": "not_interpreted", "assets_identified": len(out.assets)},
            }
            out.availability = "text_partial" if pending else "text_ready"
            out.warnings.append(
                "Native profile preserves tables/equations as page assets; use docling profile for layout structure"
            )
            return out
        finally:
            renderer.close()


class OfficeProcessor:
    def extract(self, path, mime):
        validate_zip(path)
        out = ExtractedRepresentation(
            processor_name="office",
            processor_version=version("python-docx"),
            coverage={"text": {"state": "complete"}, "visual": {"state": "not_interpreted"}},
        )
        if "wordprocessingml" in mime:
            from docx import Document

            doc = Document(path)
            for i, element in enumerate(doc.iter_inner_content()):
                loc = {"kind": "section", "value": str(i + 1)}
                if hasattr(element, "rows"):
                    rows = [[cell.text for cell in row.cells] for row in element.rows]
                    out.blocks.append(
                        Block(
                            "table",
                            "\n".join(" | ".join(row) for row in rows),
                            loc,
                            data={"rows": rows},
                        )
                    )
                else:
                    kind = "heading" if element.style.name.startswith("Heading") else "paragraph"
                    out.blocks += text_blocks(element.text, loc)
                    if out.blocks and kind == "heading":
                        out.blocks[-1].type = kind
            with zipfile.ZipFile(path) as z:
                for name in z.namelist():
                    if name.startswith("word/media/"):
                        content = z.read(name)
                        try:
                            img = Image.open(io.BytesIO(content))
                            mime_type = Image.MIME[img.format]
                        except (OSError, KeyError):
                            continue
                        a = ExtractedAsset(
                            content, mime_type, "figure", {"kind": "embedded", "value": name}
                        )
                        out.assets.append(a)
                        out.blocks.append(
                            Block(
                                "picture_ref",
                                locator=a.locator,
                                asset_refs=[f"wks://assets/{a.id}"],
                            )
                        )
        else:
            from pptx import Presentation

            doc = Presentation(path)
            for i, slide in enumerate(doc.slides, 1):
                loc = {"kind": "slide", "value": str(i)}
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        out.blocks += text_blocks(shape.text, loc)
                    if shape.has_table:
                        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                        out.blocks.append(
                            Block(
                                "table",
                                "\n".join(" | ".join(row) for row in rows),
                                loc,
                                data={"rows": rows},
                            )
                        )
                    if shape.shape_type == 13:
                        a = ExtractedAsset(
                            shape.image.blob, shape.image.content_type, "figure", loc
                        )
                        out.assets.append(a)
                        out.blocks.append(
                            Block("picture_ref", locator=loc, asset_refs=[f"wks://assets/{a.id}"])
                        )
        return out


class WebContentProcessor:
    def extract(self, path, mime):
        import trafilatura

        value = trafilatura.extract(path.read_text(errors="replace"), include_tables=True) or ""
        return ExtractedRepresentation(
            blocks=text_blocks(value),
            processor_name="trafilatura",
            processor_version=version("trafilatura"),
            availability="text_ready" if value else "metadata_only",
            coverage={"text": {"state": "complete" if value else "unavailable"}},
        )


class UnsupportedProcessor:
    def extract(self, path, mime):
        return ExtractedRepresentation(
            availability="unsupported_processing",
            processor_name="unsupported",
            coverage={"text": {"state": "unsupported"}},
        )


def processor(settings, mime):
    if mime == "application/pdf":
        return PDFProcessor(settings)
    if "openxmlformats" in mime:
        return OfficeProcessor()
    if mime == "text/html":
        return WebContentProcessor()
    if mime.startswith("text/") or mime in {"application/json", "application/yaml"}:
        return TextProcessor()
    if mime.startswith("image/"):
        return ImageProcessor(settings)
    return UnsupportedProcessor()
