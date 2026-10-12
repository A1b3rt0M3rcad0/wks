"""Offline spreadsheet, EPUB and structured-text decoders; no active content execution."""

import posixpath
import zipfile
from html.parser import HTMLParser
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree as ET

from wks_core.domain.models import Block, Error, ExtractedAsset, ExtractedRepresentation
from wks_core.storage.content import validate_zip


def xml(raw):
    lowered = raw.lower()
    if b"<!doctype" in lowered or b"<!entity" in lowered:
        raise Error("processing.content_invalid", "External entities and doctypes are disabled")
    return ET.fromstring(raw)


class DocumentHTML(HTMLParser):
    def __init__(self, locator=None):
        super().__init__(convert_charrefs=True)
        self.blocks, self.ignored, self.tag, self.text, self.locator = (
            [],
            0,
            "p",
            [],
            locator or {"kind": "html"},
        )
        self.images = []

    def handle_starttag(self, tag, attributes):
        if tag in {"script", "style", "iframe", "object"}:
            self.ignored += 1
        if self.ignored:
            return
        if tag in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "pre"}:
            self.flush()
            self.tag = tag
        if tag == "img":
            self.images.append(dict(attributes).get("src", ""))
        if tag == "br":
            self.text.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "iframe", "object"} and self.ignored:
            self.ignored -= 1
            return
        if not self.ignored and tag in {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "pre"}:
            self.flush()
            self.tag = "p"

    def handle_data(self, data):
        if not self.ignored:
            self.text.append(data)

    def flush(self):
        value = "".join(self.text).strip()
        self.text = []
        if not value:
            return
        kind = (
            "heading"
            if self.tag.startswith("h") and self.tag[1:].isdigit()
            else "code"
            if self.tag == "pre"
            else "list"
            if self.tag == "li"
            else "paragraph"
        )
        data = (
            {"level": int(self.tag[1:])}
            if kind == "heading"
            else {"items": [value]}
            if kind == "list"
            else {}
        )
        self.blocks.append(
            Block(kind, value, {**self.locator, "element": len(self.blocks)}, data=data)
        )


class StructuredTextProcessor:
    def extract(self, path, mime):
        value = path.read_text(encoding="utf-8-sig")
        if mime in {"application/xml", "text/xml"}:
            xml(value.encode())
        if mime == "text/html":
            parser = DocumentHTML()
            parser.feed(value)
            parser.close()
            parser.flush()
            return ExtractedRepresentation(
                blocks=parser.blocks,
                processor_name="html-offline",
                coverage={"text": {"state": "complete"}, "visual": {"state": "not_interpreted"}},
                warnings=["external_resources_not_fetched"] if parser.images else [],
            )
        language = {
            "application/xml": "xml",
            "text/xml": "xml",
            "application/javascript": "javascript",
            "text/javascript": "javascript",
            "text/x-python": "python",
            "text/x-shellscript": "bash",
            "text/x-c": "c",
            "text/x-c++": "cpp",
        }.get(mime, "")
        return ExtractedRepresentation(
            blocks=[
                Block(
                    "code",
                    value,
                    {"kind": "line_range", "start_line": 1, "end_line": len(value.splitlines())},
                    data={"language": language},
                )
            ],
            processor_name="structured-text",
            coverage={"text": {"state": "complete"}},
        )


class SpreadsheetProcessor:
    def __init__(self, settings):
        self.settings = settings

    def extract(self, path, mime):
        validate_zip(path)
        output = ExtractedRepresentation(
            processor_name="spreadsheet",
            coverage={
                "text": {"state": "complete"},
                "formulas": {"state": "preserved_not_executed"},
            },
        )
        limit, count = self.settings.process_max_table_rows, 0

        def append(sheet, first, rows):
            output.blocks.append(
                Block(
                    "table",
                    "\n".join(" | ".join(row) for row in rows),
                    {
                        "kind": "sheet_row_range",
                        "sheet": sheet,
                        "start_row": first,
                        "end_row": first + len(rows) - 1,
                    },
                    data={"rows": rows, "formulas_evaluated": False},
                )
            )

        if "spreadsheetml" in mime:
            import openpyxl

            stream = path.open("rb")
            workbook = openpyxl.load_workbook(
                stream, read_only=True, data_only=False, keep_links=False
            )
            try:
                for sheet in workbook:
                    if sheet.max_column and sheet.max_column > 4096:
                        raise Error(
                            "processing.content_invalid", "Spreadsheet column budget exceeded"
                        )
                    chunk, first = [], 1
                    for index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        count += 1
                        if count > limit:
                            break
                        chunk.append([str(cell) if cell is not None else "" for cell in row])
                        if len(chunk) == 20:
                            append(sheet.title, first, chunk)
                            chunk, first = [], index + 1
                    if chunk:
                        append(sheet.title, first, chunk)
                    if count > limit:
                        break
            finally:
                workbook.close()
                stream.close()
        else:
            with zipfile.ZipFile(path) as archive:
                root = xml(archive.read("content.xml"))
            ns = {
                "t": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
                "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
            }
            attribute = "{" + ns["t"] + "}"
            for sheet in root.findall(".//t:table", ns):
                chunk, first, index = [], 1, 0
                for row in sheet.findall("t:table-row", ns):
                    repeats = min(int(row.get(attribute + "number-rows-repeated", "1")), limit + 1)
                    values = []
                    for cell in row:
                        repeat = int(cell.get(attribute + "number-columns-repeated", "1"))
                        if repeat < 1 or len(values) + repeat > 4096:
                            raise Error(
                                "processing.content_invalid", "Spreadsheet column budget exceeded"
                            )
                        formula = cell.get(attribute + "formula")
                        values.extend([formula or " ".join(cell.itertext()).strip()] * repeat)
                    for _ in range(repeats):
                        index += 1
                        count += 1
                        if count > limit:
                            break
                        chunk.append(values)
                        if len(chunk) == 20:
                            append(sheet.get(attribute + "name", "Sheet"), first, chunk)
                            chunk, first = [], index + 1
                    if count > limit:
                        break
                if chunk:
                    append(sheet.get(attribute + "name", "Sheet"), first, chunk)
                if count > limit:
                    break
        if count > limit:
            output.availability = "text_partial"
            output.coverage["text"]["state"] = "partial"
            output.warnings.append("spreadsheet_row_budget_reached")
        output.coverage["quality"] = {"state": "not_measured", "rows_extracted": min(count, limit)}
        return output


class EPUBProcessor:
    def __init__(self, settings):
        self.settings = settings

    def extract(self, path, mime):
        validate_zip(path)
        output = ExtractedRepresentation(
            processor_name="epub-offline",
            coverage={"text": {"state": "complete"}, "visual": {"state": "not_interpreted"}},
        )

        def relative(base, href):
            if urlsplit(href).scheme or href.startswith("/"):
                raise Error("processing.content_invalid", "External EPUB resource is disabled")
            target = posixpath.normpath(posixpath.join(base, unquote(href).split("#")[0]))
            if target == ".." or target.startswith("../") or "\\" in target:
                raise Error("processing.content_invalid", "Invalid EPUB resource path")
            return target

        with zipfile.ZipFile(path) as archive:
            container = xml(archive.read("META-INF/container.xml"))
            rootfile = next(
                container.iter("{urn:oasis:names:tc:opendocument:xmlns:container}rootfile")
            )
            package = relative("", rootfile.get("full-path", ""))
            root = xml(archive.read(package))
            base = posixpath.dirname(package)
            manifest = {
                item.get("id"): item for item in root.iter("{http://www.idpf.org/2007/opf}item")
            }
            chapters = list(root.iter("{http://www.idpf.org/2007/opf}itemref"))
            for index, reference in enumerate(chapters[: self.settings.process_max_pages], 1):
                item = manifest[reference.get("idref")]
                target = relative(base, item.get("href", ""))
                parser = DocumentHTML({"kind": "epub_chapter", "chapter": index, "path": target})
                parser.feed(archive.read(target).decode("utf-8"))
                parser.close()
                parser.flush()
                output.blocks.extend(parser.blocks)
                for src in parser.images:
                    resource = relative(posixpath.dirname(target), src)
                    content = archive.read(resource)
                    import io

                    from PIL import Image

                    with Image.open(io.BytesIO(content)) as image:
                        image.verify()
                        asset_mime = Image.MIME[image.format]
                    asset = ExtractedAsset(
                        content,
                        asset_mime,
                        "figure",
                        {"kind": "epub_resource", "path": resource, "chapter": index},
                    )
                    output.assets.append(asset)
                    output.blocks.append(
                        Block(
                            "picture_ref",
                            locator=asset.locator,
                            asset_refs=[f"wks://assets/{asset.id}"],
                        )
                    )
            if len(chapters) > self.settings.process_max_pages:
                output.availability = "text_partial"
                output.coverage["text"]["state"] = "partial"
                output.warnings.append("chapter_budget_reached")
        output.coverage["quality"] = {
            "state": "not_measured",
            "chapters_extracted": min(len(chapters), self.settings.process_max_pages),
        }
        return output
