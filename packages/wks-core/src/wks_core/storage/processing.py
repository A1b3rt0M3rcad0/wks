"""Bounded deterministic processors. Heavy converters run in worker subprocesses."""

import csv
import io
import json
import subprocess
import zipfile
from dataclasses import asdict
from importlib.metadata import version
from pathlib import Path

from PIL import Image
from wks_core.domain.models import Block, Error, ExtractedAsset, ExtractedRepresentation


def detect_mime(path, declared="application/octet-stream"):
    with path.open("rb") as f:
        head = f.read(8192)
    signatures = [
        (b"%PDF-", "application/pdf"),
        (b"\x89PNG\r\n\x1a\n", "image/png"),
        (b"\xff\xd8\xff", "image/jpeg"),
        (b"GIF8", "image/gif"),
        (b"ID3", "audio/mpeg"),
        (b"fLaC", "audio/flac"),
        (b"OggS", "audio/ogg"),
    ]
    for signature, mime in signatures:
        if head.startswith(signature):
            return mime
    if head.startswith(b"RIFF"):
        if head[8:12] == b"WAVE":
            return "audio/wav"
        if head[8:12] == b"WEBP":
            return "image/webp"
    if head[4:8] == b"ftyp":
        return "audio/mp4" if head[8:12] in {b"M4A ", b"M4B "} else "video/mp4"
    if len(head) >= 2 and head[0] == 255 and head[1] & 0xE0 == 0xE0:
        return "audio/mpeg"
    if head.startswith(b"\x1aE\xdf\xa3"):
        return "video/webm"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
                if any(n.startswith("word/") for n in names):
                    return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                if any(n.startswith("ppt/") for n in names):
                    return (
                        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                    )
        except zipfile.BadZipFile:
            pass
    if b"\0" not in head:
        try:
            head.decode("utf-8")
            if declared in {
                "text/plain",
                "text/markdown",
                "text/csv",
                "text/html",
                "application/json",
                "application/yaml",
                "text/yaml",
            }:
                return declared
            if head and all(b >= 32 or b in {9, 10, 13} for b in head):
                return "text/plain"
        except UnicodeDecodeError:
            pass
    return "application/octet-stream"


def validate_zip(path, max_bytes=128 * 1024 * 1024):
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        if len(infos) > 10000 or sum(x.file_size for x in infos) > max_bytes:
            raise Error("upload.content_invalid", "Archive expansion exceeds limit")
        for info in infos:
            if (
                Path(info.filename).is_absolute()
                or ".." in Path(info.filename).parts
                or info.file_size > max(1024 * 1024, info.compress_size * 200)
            ):
                raise Error("upload.content_invalid", "Unsafe archive")


def text_blocks(value, locator=None, origin="native_text"):
    blocks, parent = [], None
    for line_no, line in enumerate(value.splitlines(), 1):
        if not line.strip():
            continue
        loc = locator or {"kind": "line_range", "start_line": line_no, "end_line": line_no}
        for offset in range(0, len(line), 2000):
            part = line[offset : offset + 2000]
            b = Block("paragraph", part, loc, origin_kind=origin, parent_id=parent)
            if part.startswith("#") and offset == 0:
                level = len(part) - len(part.lstrip("#"))
                b.type, b.text, b.data = "heading", part.lstrip("# "), {"level": min(level, 6)}
                parent = b.id
            blocks.append(b)
    return blocks


class TextProcessor:
    def extract(self, path, mime):
        try:
            value = path.read_text(encoding="utf-8-sig")
            if mime == "application/json":
                json.loads(value)
            elif mime in {"application/yaml", "text/yaml"}:
                import yaml

                yaml.safe_load(value)
        except (ValueError, UnicodeError, RecursionError) as exc:
            return ExtractedRepresentation(
                availability="metadata_only",
                warnings=[f"invalid_text:{type(exc).__name__}"],
                coverage={"text": {"state": "failed"}},
            )
        if mime == "text/csv":
            rows = list(csv.reader(io.StringIO(value)))
            blocks = []
            # Chunk rows instead of placing arbitrary CSV content in a single response.
            for start in range(0, len(rows), 20):
                chunk = rows[start : start + 20]
                blocks.append(
                    Block(
                        "table",
                        "\n".join(", ".join(r) for r in chunk),
                        {"kind": "row_range", "start": start + 1, "end": start + len(chunk)},
                        data={"rows": chunk},
                    )
                )
        else:
            blocks = text_blocks(value)
        return ExtractedRepresentation(
            blocks=blocks, coverage={"text": {"state": "complete"}}, processor_name="text"
        )


class ImageProcessor:
    def __init__(self, settings):
        self.settings = settings

    def ocr(self, path):
        try:
            result = subprocess.run(
                ["tesseract", str(path), "stdout", "-l", self.settings.ocr_language, "--psm", "3"],
                capture_output=True,
                timeout=self.settings.processing_timeout_seconds,
                check=True,
            )
            return result.stdout.decode().strip(), None
        except (OSError, subprocess.SubprocessError):
            return "", "ocr_provider_unavailable"

    def extract(self, path, mime):
        Image.MAX_IMAGE_PIXELS = 40_000_000
        with Image.open(path) as img:
            img.verify()
        data = path.read_bytes()
        a = ExtractedAsset(data, mime, "image", {"kind": "image", "value": "original"})
        b = Block("picture_ref", locator=a.locator, asset_refs=[f"wks://assets/{a.id}"])
        result = ExtractedRepresentation(
            blocks=[b],
            assets=[a],
            availability="metadata_only",
            processor_name="image",
            processor_version=version("pillow"),
            coverage={"text": {"state": "pending"}, "visual": {"state": "not_interpreted"}},
        )
        if self.settings.ocr_enabled:
            value, warning = self.ocr(path)
            result.blocks += text_blocks(value, a.locator, "ocr")
            for block in result.blocks:
                block.asset_refs = [f"wks://assets/{a.id}"]
            result.coverage["ocr"] = {"state": "failed" if warning else "complete"}
            result.coverage["text"] = {"state": "partial" if warning else "complete"}
            result.availability = "text_partial" if warning else "text_ready"
            if warning:
                result.warnings.append(warning)
        return result


class PDFProcessor:
    """Native text first; preserve rendered pages and embedded image bytes; selective OCR."""

    def __init__(self, settings):
        self.settings = settings

    def extract(self, path, mime):
        import pypdfium2
        from pypdf import PdfReader

        pdf = PdfReader(path)
        if pdf.is_encrypted:
            raise Error("upload.content_invalid", "Encrypted PDF cannot be processed")
        renderer = pypdfium2.PdfDocument(str(path))
        out = ExtractedRepresentation(
            processor_name="native-pdf", processor_version=version("pypdf")
        )
        pending, completed = 0, 0
        try:
            for page_no, page in enumerate(pdf.pages[: self.settings.process_max_pages], 1):
                loc = {"kind": "page", "value": str(page_no)}
                out.blocks.append(Block("heading", f"Page {page_no}", loc, data={"level": 2}))
                try:
                    value = page.extract_text() or ""
                except Exception:
                    value = ""
                    out.warnings.append(f"page_{page_no}_native_extraction_failed")
                import math

                size = renderer[page_no - 1].get_size()
                scale = min(2, math.sqrt(16_000_000 / max(1, size[0] * size[1])))
                page_image = renderer[page_no - 1].render(scale=scale).to_pil()
                bio = io.BytesIO()
                page_image.save(bio, format="PNG")
                page_asset = ExtractedAsset(bio.getvalue(), "image/png", "page", loc)
                out.assets.append(page_asset)
                page_refs = [f"wks://assets/{page_asset.id}"]
                try:
                    for picture in page.images:
                        img = Image.open(io.BytesIO(picture.data))
                        media_type = Image.MIME.get(img.format, "application/octet-stream")
                        a = ExtractedAsset(picture.data, media_type, "figure", loc)
                        out.assets.append(a)
                        page_refs.append(f"wks://assets/{a.id}")
                except Exception:
                    out.warnings.append(f"page_{page_no}_image_preserved_as_page")
                origin = "native_text"
                if not value.strip() and self.settings.ocr_enabled:
                    from tempfile import NamedTemporaryFile

                    with NamedTemporaryFile(suffix=".png") as f:
                        f.write(page_asset.content)
                        f.flush()
                        value, warning = ImageProcessor(self.settings).ocr(Path(f.name))
                    origin = "ocr"
                    if warning:
                        out.warnings.append(f"page_{page_no}_{warning}")
                if value.strip():
                    completed += 1
                    blocks = text_blocks(value, loc, origin)
                    for block in blocks:
                        block.asset_refs = page_refs
                    out.blocks += blocks
                else:
                    pending += 1
                out.blocks.append(Block("picture_ref", locator=loc, asset_refs=page_refs))
            pending += max(0, len(pdf.pages) - self.settings.process_max_pages)
            out.coverage = {
                "text": {
                    "state": "partial" if pending else "complete",
                    "processed_pages": completed,
                    "total_pages": len(pdf.pages),
                    "pending_pages": pending,
                },
                "visual": {"state": "not_interpreted", "assets_identified": len(out.assets)},
            }
            out.availability = "text_partial" if pending else "text_ready"
            out.warnings.append(
                "Native profile preserves tables/equations as page assets; use docling profile for layout structure"
            )
            return out
        finally:
            renderer.close()


class OfficeProcessor:
    def extract(self, path, mime):
        validate_zip(path)
        out = ExtractedRepresentation(
            processor_name="office",
            processor_version=version("python-docx"),
            coverage={"text": {"state": "complete"}, "visual": {"state": "not_interpreted"}},
        )
        if "wordprocessingml" in mime:
            from docx import Document

            doc = Document(path)
            for i, element in enumerate(doc.iter_inner_content()):
                loc = {"kind": "section", "value": str(i + 1)}
                if hasattr(element, "rows"):
                    rows = [[cell.text for cell in row.cells] for row in element.rows]
                    out.blocks.append(
                        Block(
                            "table",
                            "\n".join(" | ".join(row) for row in rows),
                            loc,
                            data={"rows": rows},
                        )
                    )
                else:
                    kind = "heading" if element.style.name.startswith("Heading") else "paragraph"
                    out.blocks += text_blocks(element.text, loc)
                    if out.blocks and kind == "heading":
                        out.blocks[-1].type = kind
            with zipfile.ZipFile(path) as z:
                for name in z.namelist():
                    if name.startswith("word/media/"):
                        content = z.read(name)
                        try:
                            img = Image.open(io.BytesIO(content))
                            mime_type = Image.MIME[img.format]
                        except (OSError, KeyError):
                            continue
                        a = ExtractedAsset(
                            content, mime_type, "figure", {"kind": "embedded", "value": name}
                        )
                        out.assets.append(a)
                        out.blocks.append(
                            Block(
                                "picture_ref",
                                locator=a.locator,
                                asset_refs=[f"wks://assets/{a.id}"],
                            )
                        )
        else:
            from pptx import Presentation

            doc = Presentation(path)
            for i, slide in enumerate(doc.slides, 1):
                loc = {"kind": "slide", "value": str(i)}
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        out.blocks += text_blocks(shape.text, loc)
                    if shape.has_table:
                        rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                        out.blocks.append(
                            Block(
                                "table",
                                "\n".join(" | ".join(row) for row in rows),
                                loc,
                                data={"rows": rows},
                            )
                        )
                    if shape.shape_type == 13:
                        a = ExtractedAsset(
                            shape.image.blob, shape.image.content_type, "figure", loc
                        )
                        out.assets.append(a)
                        out.blocks.append(
                            Block("picture_ref", locator=loc, asset_refs=[f"wks://assets/{a.id}"])
                        )
        return out


class DoclingProcessor:
    def __init__(self, settings):
        self.settings = settings

    def extract(self, path, mime):
        from docling.datamodel.base_models import InputFormat
        from docling.datamodel.pipeline_options import PdfPipelineOptions, TesseractCliOcrOptions
        from docling.document_converter import DocumentConverter, PdfFormatOption

        if "openxmlformats" in mime:
            validate_zip(path)
        options = PdfPipelineOptions()
        if self.settings.docling_artifacts_path:
            options.artifacts_path = Path(self.settings.docling_artifacts_path)
        options.do_ocr = self.settings.ocr_enabled
        options.ocr_options = TesseractCliOcrOptions(lang=self.settings.ocr_language.split("+"))
        options.generate_picture_images = True
        options.generate_page_images = True
        converter = DocumentConverter(
            format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)}
        )
        suffix = {"application/pdf": ".pdf"}
        if "wordprocessingml" in mime:
            suffix[mime] = ".docx"
        if "presentationml" in mime:
            suffix[mime] = ".pptx"
        import shutil
        from tempfile import TemporaryDirectory

        with TemporaryDirectory() as temp:
            source = Path(temp) / ("input" + suffix.get(mime, path.suffix))
            shutil.copyfile(path, source)
            converted = converter.convert(source, max_num_pages=self.settings.process_max_pages)
        doc = converted.document
        out = ExtractedRepresentation(
            processor_name="docling", processor_version=version("docling")
        )
        native_pages = {}
        if mime == "application/pdf":
            from pypdf import PdfReader

            native_pages = {
                i + 1: " ".join((page.extract_text() or "").split())
                for i, page in enumerate(PdfReader(path).pages)
            }
            # Page renders preserve evidence even when layout detection finds no picture.
            for number, page in doc.pages.items():
                if page.image:
                    buf = io.BytesIO()
                    page.image.pil_image.save(buf, "PNG")
                    out.assets.append(
                        ExtractedAsset(
                            buf.getvalue(),
                            "image/png",
                            "page",
                            {"kind": "page", "value": str(number)},
                        )
                    )
        for item, level in doc.iterate_items():
            provenance = item.prov[0] if getattr(item, "prov", None) else None
            loc = (
                {"kind": "page", "value": str(provenance.page_no)}
                if provenance
                else {"kind": "section", "value": str(len(out.blocks) + 1)}
            )
            if provenance and provenance.bbox:
                loc["bbox"] = provenance.bbox.model_dump(mode="json")
            label = str(item.label.value)
            value = getattr(item, "text", "")
            origin = "native_text"
            if (
                provenance
                and native_pages
                and (
                    not native_pages.get(provenance.page_no)
                    or (value and " ".join(value.split()) not in native_pages[provenance.page_no])
                )
            ):
                origin = "ocr"
            if label == "table":
                data = item.data.model_dump(mode="json")
                out.blocks.append(
                    Block(
                        "table",
                        item.export_to_markdown(doc=doc),
                        loc,
                        origin_kind=origin,
                        data=data,
                    )
                )
            elif label == "picture":
                image = item.get_image(doc)
                if image:
                    buf = io.BytesIO()
                    image.save(buf, "PNG")
                    a = ExtractedAsset(buf.getvalue(), "image/png", "figure", loc)
                    out.assets.append(a)
                    out.blocks.append(
                        Block("picture_ref", locator=loc, asset_refs=[f"wks://assets/{a.id}"])
                    )
            elif value:
                blocks = text_blocks(value, loc, origin)
                if label in {"section_header", "title"}:
                    for b in blocks:
                        b.type, b.data = "heading", {"level": min(level + 1, 6)}
                elif label == "formula":
                    for b in blocks:
                        b.type = "equation"
                out.blocks += blocks
        status = str(converted.status.value)
        out.availability = "text_ready" if status == "success" else "text_partial"
        out.coverage = {
            "text": {"state": "complete" if status == "success" else "partial"},
            "visual": {"state": "not_interpreted", "assets_identified": len(out.assets)},
        }
        out.warnings = [str(e.error_message)[:200] for e in converted.errors]
        # Associate same-page figures with text so search discovers media provenance.
        for block in out.blocks:
            block.asset_refs += [
                f"wks://assets/{a.id}"
                for a in out.assets
                if a.locator.get("value") == block.locator.get("value")
                and f"wks://assets/{a.id}" not in block.asset_refs
            ]
        return out


class DoclingWithFallback:
    def __init__(self, settings):
        self.settings = settings

    def extract(self, path, mime):
        try:
            return DoclingProcessor(self.settings).extract(path, mime)
        except Exception:
            # Model downloads/provider outages cannot erase deterministic native text.
            fallback = (
                PDFProcessor(self.settings) if mime == "application/pdf" else OfficeProcessor()
            )
            out = fallback.extract(path, mime)
            out.warnings.append("docling_unavailable_native_fallback_used")
            out.coverage["layout"] = {"state": "pending"}
            return out


class AudioProcessor:
    def __init__(self, settings):
        self.settings = settings

    def extract(self, path, mime):
        if not self.settings.asr_enabled or not self.settings.asr_model_path:
            return ExtractedRepresentation(
                availability="metadata_only",
                processor_name="audio",
                coverage={"audio": {"state": "pending"}},
                warnings=["asr_provider_unavailable"],
            )
        info = probe(path)
        duration = float(info["format"].get("duration", 0))
        if not duration or duration > self.settings.audio_max_seconds:
            raise Error("processing.duration_limit", "Audio duration outside configured limit")
        from faster_whisper import WhisperModel

        model = WhisperModel(
            self.settings.asr_model_path, device="cpu", compute_type="int8", local_files_only=True
        )
        segments, facts = model.transcribe(str(path), vad_filter=True)
        blocks = []
        for s in segments:
            loc = {
                "kind": "time_range",
                "start_ms": round(s.start * 1000),
                "end_ms": round(s.end * 1000),
            }
            blocks.append(
                Block(
                    "audio_transcript",
                    s.text.strip(),
                    loc,
                    origin_kind="transcript",
                    data={
                        "language": facts.language,
                        "avg_logprob": s.avg_logprob,
                        "no_speech_probability": s.no_speech_prob,
                        "quality": "uncertain"
                        if s.avg_logprob < -1
                        else "unverified_transcription",
                    },
                )
            )
        return ExtractedRepresentation(
            blocks=blocks,
            processor_name="faster-whisper",
            processor_version=version("faster-whisper"),
            warnings=["Transcription is probabilistic; non-speech audio is not interpreted"],
            coverage={
                "audio": {"state": "complete", "duration_seconds": duration},
                "text": {"state": "complete"},
            },
        )


def probe(path):
    p = subprocess.run(
        ["ffprobe", "-v", "error", "-show_format", "-show_streams", "-of", "json", str(path)],
        capture_output=True,
        timeout=30,
        check=True,
    )
    return json.loads(p.stdout)


class VideoProcessor:
    def __init__(self, settings):
        self.settings = settings

    def extract(self, path, mime):
        if not self.settings.video_enabled:
            return ExtractedRepresentation(
                availability="metadata_only",
                processor_name="video",
                coverage={"visual": {"state": "pending"}, "audio": {"state": "pending"}},
            )
        info = probe(path)
        duration = float(info["format"].get("duration", 0))
        if not duration or duration > self.settings.video_max_seconds:
            raise Error("processing.duration_limit", "Video duration outside configured limit")
        out = ExtractedRepresentation(
            processor_name="ffmpeg-video",
            availability="text_partial",
            coverage={"visual": {"state": "sampled", "duration_seconds": duration}},
        )
        from tempfile import TemporaryDirectory

        with TemporaryDirectory() as temp:
            for i in range(self.settings.video_max_frames):
                timestamp = duration * i / self.settings.video_max_frames
                target = Path(temp) / f"frame-{i}.png"
                subprocess.run(
                    [
                        "ffmpeg",
                        "-v",
                        "error",
                        "-ss",
                        str(timestamp),
                        "-i",
                        str(path),
                        "-frames:v",
                        "1",
                        "-threads",
                        "1",
                        str(target),
                    ],
                    timeout=30,
                    check=True,
                )
                if not target.exists():
                    continue
                loc = {"kind": "frame", "time_ms": round(timestamp * 1000), "sample_index": i}
                a = ExtractedAsset(target.read_bytes(), "image/png", "frame", loc)
                out.assets.append(a)
                out.blocks.append(
                    Block("video_frame_ref", locator=loc, asset_refs=[f"wks://assets/{a.id}"])
                )
            if any(s["codec_type"] == "audio" for s in info["streams"]):
                audio = Path(temp) / "audio.wav"
                subprocess.run(
                    [
                        "ffmpeg",
                        "-v",
                        "error",
                        "-i",
                        str(path),
                        "-vn",
                        "-ac",
                        "1",
                        "-ar",
                        "16000",
                        str(audio),
                    ],
                    timeout=60,
                    check=True,
                )
                result = AudioProcessor(self.settings).extract(audio, "audio/wav")
                out.blocks += result.blocks
                out.coverage.update(result.coverage)
                out.warnings += result.warnings
            else:
                out.coverage["audio"] = {"state": "not_applicable"}
        out.coverage["visual"]["frames_sampled"] = len(out.assets)
        return out


class WebContentProcessor:
    def extract(self, path, mime):
        import trafilatura

        value = trafilatura.extract(path.read_text(errors="replace"), include_tables=True) or ""
        return ExtractedRepresentation(
            blocks=text_blocks(value),
            processor_name="trafilatura",
            processor_version=version("trafilatura"),
            availability="text_ready" if value else "metadata_only",
            coverage={"text": {"state": "complete" if value else "unavailable"}},
        )


class UnsupportedProcessor:
    def extract(self, path, mime):
        return ExtractedRepresentation(
            availability="unsupported_processing",
            processor_name="unsupported",
            coverage={"text": {"state": "unsupported"}},
        )


def processor(settings, mime):
    if mime == "application/pdf":
        return (
            DoclingWithFallback(settings)
            if settings.extraction_profile == "docling"
            else PDFProcessor(settings)
        )
    if "openxmlformats" in mime:
        return (
            DoclingWithFallback(settings)
            if settings.extraction_profile == "docling"
            else OfficeProcessor()
        )
    if mime == "text/html":
        return WebContentProcessor()
    if mime.startswith("text/") or mime in {"application/json", "application/yaml"}:
        return TextProcessor()
    if mime.startswith("image/"):
        return ImageProcessor(settings)
    if mime.startswith("audio/"):
        return AudioProcessor(settings)
    if mime.startswith("video/"):
        return VideoProcessor(settings)
    return UnsupportedProcessor()


def main():
    import sys

    from wks_core.settings import Settings

    path, mime, config_path, destination = sys.argv[1:]
    config = json.loads(Path(config_path).read_text())
    settings = Settings(**config)
    import resource

    memory_limit = settings.process_max_memory_mb * 1024 * 1024
    resource.setrlimit(resource.RLIMIT_AS, (memory_limit, memory_limit))
    cpu_limit = settings.processing_timeout_seconds + 5
    resource.setrlimit(resource.RLIMIT_CPU, (cpu_limit, cpu_limit + 5))
    try:
        result = processor(settings, mime).extract(Path(path), mime)
        data = asdict(result)
        for asset in data["assets"]:
            target = Path(destination) / asset["id"]
            target.write_bytes(asset.pop("content"))
        (Path(destination) / "result.json").write_text(json.dumps(data, ensure_ascii=False))
    except Exception as exc:
        code = exc.code if isinstance(exc, Error) else "processing.content_invalid"
        (Path(destination) / "error.json").write_text(json.dumps({"code": code}))
        raise


if __name__ == "__main__":
    main()
